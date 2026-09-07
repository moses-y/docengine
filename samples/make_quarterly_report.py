"""Builds samples/quarterly-report.docx for the import demo.

Deliberately includes a table and an image alongside the supported
constructs, so the import path exercises both what it preserves and what it
reports as dropped.
"""
import io
import struct
import zlib
from pathlib import Path

from docx import Document
from docx.shared import Inches

doc = Document()

doc.add_heading("Quarterly Product Review", level=1)

p = doc.add_paragraph()
p.add_run("Prepared by the product team. ")
p.add_run("Confidential").bold = True
p.add_run(" — for internal circulation only, and ")
p.add_run("not final").italic = True
p.add_run(".")

doc.add_heading("Highlights", level=2)
for line in [
    "Editor rewritten on a whitelisted node schema",
    "Import supports .txt, .md and .docx",
    "Sharing with per-document viewer and editor roles",
]:
    doc.add_paragraph(line, style="List Bullet")

doc.add_heading("Next steps", level=2)
for line in [
    "Real-time co-editing with presence",
    "Version history panel",
    "Full-text search",
]:
    doc.add_paragraph(line, style="List Number")

doc.add_heading("Open questions", level=3)
doc.add_paragraph(
    "Whether to ship export before search. The table below is included on "
    "purpose: it has no whitelisted target node, so the importer should drop "
    "it and say so."
)

table = doc.add_table(rows=3, cols=2)
table.style = "Table Grid"
cells = [("Item", "Owner"), ("Co-editing", "Backend"), ("History UI", "Frontend")]
for row, (a, b) in zip(table.rows, cells):
    row.cells[0].text = a
    row.cells[1].text = b

doc.add_paragraph("A trailing paragraph after the table, to prove the parser recovers.")


def tiny_png() -> bytes:
    """A 4x4 solid PNG, built inline to avoid shipping a binary blob."""
    w = h = 4
    raw = b"".join(b"\x00" + bytes([200, 80, 40] * w) for _ in range(h))

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


doc.add_paragraph("An embedded image follows, which is also expected to be dropped:")
doc.add_picture(io.BytesIO(tiny_png()), width=Inches(0.5))

out = Path(__file__).with_name("quarterly-report.docx")
doc.save(out)
print(f"wrote {out}")
